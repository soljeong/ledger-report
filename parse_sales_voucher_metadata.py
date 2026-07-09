#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "private_sources" / "sales_voucher_groups_principal_input.xlsx"
DEFAULT_OUTPUT = BASE_DIR / "private_intermediate" / "sales_voucher_metadata.json"

FIELD_LABELS = {
    "voucher_key": "전표키",
    "date": "일자",
    "voucher": "전표",
    "company": "거래처",
    "row_count": "행수",
    "quantity": "수량",
    "supply_amount": "공급금액",
    "vat": "부가세",
    "total_amount": "합계금액",
    "product_id": "등록번호",
    "principal": "원청",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Parse sales voucher principal metadata Excel to JSON.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def normalize(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_principal(value: Any) -> tuple[str | None, str | None]:
    raw = None if value is None else str(value)
    normalized = normalize(value)
    return raw, normalized or None


def number_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else value
    if isinstance(value, str):
        stripped = value.strip().replace(",", "")
        if stripped == "":
            return None
        try:
            parsed = float(stripped)
        except ValueError:
            return value.strip()
        return int(parsed) if parsed.is_integer() else parsed
    return value


def find_header_map(row: tuple[Any, ...]) -> dict[str, int]:
    label_to_field = {label: field for field, label in FIELD_LABELS.items()}
    header_map: dict[str, int] = {}
    for index, value in enumerate(row):
        field = label_to_field.get(normalize(value))
        if field:
            header_map[field] = index
    required = ["voucher_key", "date", "voucher", "total_amount", "principal"]
    missing = [field for field in required if field not in header_map]
    if missing:
        raise ValueError(f"missing required headers: {missing}")
    return header_map


def cell(row: tuple[Any, ...], header_map: dict[str, int], field: str) -> Any:
    index = header_map.get(field)
    if index is None or index >= len(row):
        return None
    return row[index]


def parse_metadata_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True)
    if "매출_전표별" not in workbook.sheetnames:
        raise ValueError(f"missing sheet '매출_전표별': {path}")
    sheet = workbook["매출_전표별"]
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    header_map = find_header_map(header_row)

    records: list[dict[str, Any]] = []
    for excel_row, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        voucher_key = normalize(cell(row, header_map, "voucher_key"))
        if not voucher_key:
            continue
        principal_raw, principal = normalize_principal(cell(row, header_map, "principal"))
        records.append(
            {
                "excel_row": excel_row,
                "voucher_key": voucher_key,
                "date": normalize(cell(row, header_map, "date")),
                "voucher": number_value(cell(row, header_map, "voucher")),
                "company": normalize(cell(row, header_map, "company")) or None,
                "row_count": number_value(cell(row, header_map, "row_count")),
                "quantity": number_value(cell(row, header_map, "quantity")),
                "supply_amount": number_value(cell(row, header_map, "supply_amount")),
                "vat": number_value(cell(row, header_map, "vat")),
                "total_amount": number_value(cell(row, header_map, "total_amount")),
                "product_id": number_value(cell(row, header_map, "product_id")),
                "principal_raw": principal_raw,
                "principal": principal,
            }
        )

    payload = {
        "metadata": {
            "source_file": path.name,
            "sheet_name": sheet.title,
            "record_count": len(records),
            "unique_principal_count": len({row["principal"] for row in records if row["principal"]}),
            "total_amount": sum(row["total_amount"] or 0 for row in records),
        },
        "records": records,
    }
    return payload


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> int:
    args = parse_args()
    payload = parse_metadata_workbook(args.input)
    write_json(args.output, payload)
    meta = payload["metadata"]
    print(
        f"{args.output.name}: records={meta['record_count']} "
        f"principals={meta['unique_principal_count']} total={meta['total_amount']}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
