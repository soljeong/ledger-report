#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = BASE_DIR / "private_intermediate"
DEFAULT_INPUT = DEFAULT_DATA_DIR / "validation.json"
DEFAULT_OUTPUT = DEFAULT_DATA_DIR / "selected_validation.xlsx"


HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(bold=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export selected validation buckets to Excel.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def load_validation(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def money(value: Any) -> Any:
    return value


def join_values(values: list[Any]) -> str:
    return "\n".join(str(value) for value in values if value not in (None, ""))


def row_list(rows: list[dict[str, Any]]) -> str:
    return "\n".join(
        f"{row['type']} row {row['excel_row']} {row['date']} {row['company']} "
        f"qty={row['quantity']} unit={row['unit_price']} total={row['total_amount']}"
        for row in rows
    )


def setup_sheet(ws, headers: list[str]) -> None:
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def append_rows(ws, rows: list[list[Any]]) -> None:
    for row in rows:
        ws.append(row)
    for worksheet_row in ws.iter_rows(min_row=2):
        for cell in worksheet_row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def auto_size(ws) -> None:
    for column_cells in ws.columns:
        width = 10
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            longest_line = max(len(line) for line in str(value).splitlines())
            width = max(width, min(longest_line + 2, 60))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def export_summary(wb: Workbook, validation: dict[str, Any]) -> None:
    summary = validation["summary"]
    ws = wb.active
    ws.title = "요약"
    setup_sheet(ws, ["항목", "값"])
    append_rows(
        ws,
        [
            ["같은 품명인데 규격이 다른 품명", summary["same_item_different_specs_count"]],
            ["매입 없이 매출만 있는 품명", summary["sales_without_purchase_item_count"]],
            [
                "매입 품명은 있으나 같은 품명+규격 매입이 없는 매출 품목",
                summary["sales_without_purchase_item_spec_count"],
            ],
            ["출력 파일", DEFAULT_OUTPUT.name],
        ],
    )
    auto_size(ws)


def export_same_item_different_specs(wb: Workbook, entries: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("같은품명_다른규격")
    setup_sheet(
        ws,
        [
            "품명 key",
            "표시 품명",
            "규격 수",
            "매입 행수",
            "매출 행수",
            "규격 목록",
            "관련 원본 행",
        ],
    )
    rows = []
    for entry in entries:
        all_rows = [row for spec in entry["specifications"] for row in spec["rows"]]
        specs = [spec_text for spec in entry["specifications"] for spec_text in spec["specifications"]]
        rows.append(
            [
                entry["item_key"],
                join_values(entry["item_names"]),
                len(entry["specifications"]),
                entry["side_counts"].get("purchase", 0),
                entry["side_counts"].get("sales", 0),
                join_values(specs),
                row_list(all_rows),
            ]
        )
    append_rows(ws, rows)
    auto_size(ws)


def export_sales_without_purchase_item(wb: Workbook, entries: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("매출품명_매입없음")
    setup_sheet(
        ws,
        [
            "품명 key",
            "표시 품명",
            "매출 행수",
            "매출 합계금액",
            "규격 목록",
            "관련 매출 행",
        ],
    )
    append_rows(
        ws,
        [
            [
                entry["item_key"],
                join_values(entry["item_names"]),
                entry["sales_rows"],
                money(entry["sales_total_amount"]),
                join_values(entry["specifications"]),
                row_list(entry["rows"]),
            ]
            for entry in entries
        ],
    )
    auto_size(ws)


def export_sales_without_purchase_item_spec(wb: Workbook, entries: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("매출품명규격_매입없음")
    setup_sheet(
        ws,
        [
            "품명 key",
            "규격 key",
            "표시 품명",
            "표시 규격",
            "매출 행수",
            "매출 합계금액",
            "관련 매출 행",
        ],
    )
    append_rows(
        ws,
        [
            [
                entry["item_key"],
                entry["spec_key"],
                join_values(entry["item_names"]),
                join_values(entry["specifications"]),
                entry["sales_rows"],
                money(entry["sales_total_amount"]),
                row_list(entry["rows"]),
            ]
            for entry in entries
        ],
    )
    auto_size(ws)


def export_workbook(validation: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    export_summary(wb, validation)
    export_same_item_different_specs(wb, validation["same_item_different_specs"])
    export_sales_without_purchase_item(wb, validation["sales_without_purchase_item"])
    export_sales_without_purchase_item_spec(wb, validation["sales_without_purchase_item_spec"])
    wb.save(output_path)


def main() -> int:
    args = parse_args()
    validation = load_validation(args.input)
    export_workbook(validation, args.output)
    print(f"wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
