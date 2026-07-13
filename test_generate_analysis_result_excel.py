from __future__ import annotations

import copy
import json
import re
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from analyze_inventory import build_reconciliation
from generate_analysis_report_html_core import build_report_data, load_sources
from generate_analysis_result_excel import (
    FIFO_CALCULATION_COLUMNS,
    SHEET_NAMES,
    build_fifo_calculation_rows,
    export_workbook,
    generate_report,
    load_analysis_sources,
)


BASE_DIR = Path(__file__).resolve().parent
EXAMPLE_DIR = BASE_DIR / "examples" / "dummy_json"


class GenerateAnalysisResultExcelTests(unittest.TestCase):
    def _create_workbook(self, sources=None):
        tempdir = tempfile.TemporaryDirectory()
        output = Path(tempdir.name) / "analysis_result.xlsx"
        export_workbook(sources or load_sources(EXAMPLE_DIR), output)
        self.addCleanup(tempdir.cleanup)
        return output, load_workbook(output, data_only=False)

    @staticmethod
    def _fifo_report_data(purchases, sales, inventory_quantity):
        reconciliation = build_reconciliation(
            {"records": purchases},
            {"records": sales},
            {
                "records": [
                    {
                        "product_id": 1,
                        "item_name": "Unallocated Item",
                        "specification": "A",
                        "stock_quantity": inventory_quantity,
                    }
                ]
            },
            "2026-01-01",
            "2026-01-31",
            "2026-01-31",
        )
        return {"reconciliation": reconciliation, "records": {"purchase": purchases, "sales": sales}}

    @staticmethod
    def _sale_record(date, quantity, voucher, excel_row):
        amount = quantity * 100
        return {
            "date": date,
            "voucher": voucher,
            "excel_row": excel_row,
            "product_id": 1,
            "item_name": "Unallocated Item",
            "specification": "A",
            "quantity": quantity,
            "unit_price": 100,
            "supply_amount": amount,
            "vat": 0,
            "total_amount": amount,
        }

    @staticmethod
    def _labeled_row(workbook, sheet_name: str, label: str):
        for row in workbook[sheet_name].iter_rows(min_col=1, max_col=3):
            if row[0].value == label:
                return row
        raise AssertionError(f"{label!r} was not found in {sheet_name}")

    @staticmethod
    def _header_column(workbook, sheet_name: str, header: str):
        for row in workbook[sheet_name].iter_rows():
            for cell in row:
                if cell.value == header:
                    return cell.row, cell.column
        raise AssertionError(f"{header!r} header was not found in {sheet_name}")

    def test_creates_reopenable_workbook_with_required_sheet_order_and_json_record_counts(self):
        output, workbook = self._create_workbook()

        self.assertTrue(output.exists())
        self.assertEqual(workbook.sheetnames, SHEET_NAMES)
        sources = load_sources(EXAMPLE_DIR)
        for sheet_name, source_name in (("매입", "purchase"), ("매출", "sales"), ("재고", "inventory")):
            self.assertEqual(workbook[sheet_name].max_row - 1, len(sources[source_name]["records"]))
            self.assertEqual(workbook[sheet_name].freeze_panes, "A2")
            self.assertIsNotNone(workbook[sheet_name].auto_filter.ref)

    def test_kpis_come_from_common_report_data_and_warning_status_is_explicit(self):
        output, workbook = self._create_workbook()
        report_data = build_report_data(load_sources(EXAMPLE_DIR))
        summary_values = {
            row[0].value: row[1].value
            for row in workbook["요약"].iter_rows(min_row=2, max_col=3)
        }

        self.assertEqual(summary_values["분석 결과 상태"], "completed_with_warnings")
        self.assertEqual(summary_values["매출 공급가액"], report_data["kpis"]["sales_supply_amount"])
        self.assertEqual(summary_values["매출원가"], report_data["kpis"]["fifo_sales_cost_amount"])
        self.assertEqual(summary_values["매출총이익"], report_data["kpis"]["gross_profit"])
        self.assertIsNone(summary_values["매출총이익률"])
        self.assertEqual(workbook["손익분석"]["A1"].value, "손익 요약")
        self.assertTrue(output.exists())

    def test_fifo_calculation_sheet_uses_the_required_structure_without_changing_charts(self):
        output, workbook = self._create_workbook()
        worksheet = workbook["FIFO계산"]

        self.assertEqual(workbook.sheetnames.index("FIFO계산"), workbook.sheetnames.index("재고분석") + 1)
        self.assertEqual(workbook.sheetnames.index("FIFO계산"), workbook.sheetnames.index("매출분석") - 1)
        self.assertEqual([cell.value for cell in worksheet[1]], FIFO_CALCULATION_COLUMNS)
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, f"A1:R{worksheet.max_row}")
        self.assertEqual(sum(len(sheet._charts) for sheet in workbook.worksheets), 6)
        self.assertFalse(
            {"cost_status", "profit_status", "gross_profit_status", "amount_validation_status", "quantity_reconciliation_status"}
            & set(FIFO_CALCULATION_COLUMNS)
        )
        self.assertTrue(output.exists())

    def test_fifo_calculation_rows_flatten_existing_events_in_fifo_order(self):
        sources = load_sources(EXAMPLE_DIR)
        original_sources = copy.deepcopy(sources)
        report_data = build_report_data(sources)
        rows = build_fifo_calculation_rows(report_data)
        reconciliation = sources["reconciliation"]

        self.assertEqual(sources, original_sources)
        self.assertEqual([set(row) for row in rows], [set(FIFO_CALCULATION_COLUMNS)] * len(rows))

        completed_products = []
        previous_product_id = object()
        for row in rows:
            if row["product_id"] != previous_product_id:
                self.assertNotIn(row["product_id"], completed_products)
                completed_products.append(row["product_id"])
                previous_product_id = row["product_id"]
        self.assertEqual(completed_products, sorted(completed_products, key=str))

        same_day_rows = [
            row for row in rows
            if row["product_id"] == 1001 and row["transaction_date"] == "2026-05-10"
        ]
        self.assertEqual([row["transaction_type"] for row in same_day_rows], ["purchase", "sales", "sales"])
        opening_purchase = next(row for row in rows if row["transaction_type"] == "purchase" and row["transaction_id"] == "2026-05-01|1|2")
        self.assertEqual(opening_purchase["source_purchase_unit_cost"], 100)
        self.assertIsNone(opening_purchase["fifo_cost_amount"])
        self.assertIsNone(opening_purchase["allocated_quantity"])

        layered_sale = [
            row for row in rows
            if row["transaction_type"] == "sales" and row["transaction_id"] == "2026-05-10|2|3"
        ]
        self.assertEqual([row["allocation_sequence"] for row in layered_sale], [1, 2])
        self.assertEqual([row["allocated_quantity"] for row in layered_sale], [6, 2])
        self.assertEqual(
            [
                (row["source_purchase_date"], row["source_purchase_voucher"], row["source_purchase_excel_row"], row["source_purchase_unit_cost"], row["fifo_cost_amount"])
                for row in layered_sale
            ],
            [("2026-05-01", 1, 2, 100, 600), ("2026-05-10", 2, 3, 120, 240)],
        )

        backfill_rows = [row for row in rows if row["transaction_id"] == "2026-05-12|4|5" and row["transaction_type"] == "sales"]
        self.assertEqual([row["allocation_type"] for row in backfill_rows], ["backfill", "backfill"])
        self.assertEqual([row["source_purchase_date"] for row in backfill_rows], ["2026-05-21", "2026-05-22"])
        self.assertEqual([row["allocated_quantity"] for row in backfill_rows], [6, 4])
        self.assertEqual([row["fifo_cost_amount"] for row in backfill_rows], [1200, 880])
        self.assertTrue(any(row["transaction_type"] == "purchase" and row["transaction_id"] == "2026-05-21|4|5" for row in rows))
        self.assertFalse(any(row["product_id"] == 1004 for row in rows))

        unallocated_rows = [row for row in rows if row["allocation_type"] == "unallocated"]
        self.assertEqual(
            unallocated_rows,
            [
                {
                    **{column: None for column in FIFO_CALCULATION_COLUMNS},
                    "product_id": 1007,
                    "item_name": "Unconfirmed Part",
                    "specification": "G",
                    "transaction_date": "2026-05-14",
                    "transaction_type": "sales",
                    "transaction_id": "2026-05-14|10|11",
                    "voucher": 10,
                    "excel_row": 11,
                    "transaction_quantity": 5,
                    "allocation_sequence": 1,
                    "allocation_type": "unallocated",
                    "unallocated_quantity": 5,
                }
            ],
        )

        purchase_cancellation = [row for row in rows if row["transaction_type"] == "purchase_cancellation"]
        sales_cancellation = [row for row in rows if row["transaction_type"] == "sales_cancellation"]
        self.assertEqual([(row["allocated_quantity"], row["fifo_cost_amount"]) for row in purchase_cancellation], [(-3, -180)])
        self.assertEqual([(row["allocated_quantity"], row["fifo_cost_amount"]) for row in sales_cancellation], [(-3, -240), (-2, -160)])

        cancellation_first_sources = copy.deepcopy(sources)
        purchase_events = cancellation_first_sources["reconciliation"]["purchase_cost_events"]
        cancellation_event = next(event for event in purchase_events if event["event_type"] == "purchase_cancellation")
        purchase_events.remove(cancellation_event)
        purchase_events.insert(0, cancellation_event)
        cancellation_first_rows = build_fifo_calculation_rows(build_report_data(cancellation_first_sources))
        self.assertTrue(any(row["transaction_type"] == "purchase_cancellation" for row in cancellation_first_rows))

        period_start = reconciliation["metadata"]["period_start"]
        period_end = reconciliation["metadata"]["period_end"]
        displayed_sales_cost = sum(
            row["fifo_cost_amount"] or 0
            for row in rows
            if row["transaction_type"] in {"sales", "sales_cancellation"}
            and period_start <= row["transaction_date"] <= period_end
        )
        published_sales_cost = sum(
            event["cost_amount"]
            for event in reconciliation["sales_cost_events"]
            if period_start <= event["date"] <= period_end
        )
        self.assertEqual(displayed_sales_cost, published_sales_cost)
        self.assertEqual(displayed_sales_cost, reconciliation["summary"]["fifo_sales_cost_amount"])

    def test_fifo_cancellation_reference_keeps_purchase_and_sales_streams_separate(self):
        sources = copy.deepcopy(load_sources(EXAMPLE_DIR))
        reconciliation = sources["reconciliation"]
        purchase_cancellation = next(event for event in reconciliation["cancellation_events"] if event["type"] == "purchase_cancellation")
        sales_cancellation = next(event for event in reconciliation["cancellation_events"] if event["type"] == "sales_cancellation")
        sales_cost_event = next(event for event in reconciliation["sales_cost_events"] if event["event_type"] == "sale_cancellation")
        sales_source = next(record for record in sources["sales"]["records"] if record["quantity"] < 0)
        shared_reference = {
            "date": purchase_cancellation["date"],
            "voucher": purchase_cancellation["voucher"],
            "excel_row": purchase_cancellation["excel_row"],
        }
        for event in (sales_cancellation, sales_cost_event, sales_source):
            event.update(shared_reference)

        rows = build_fifo_calculation_rows(build_report_data(sources))
        purchase_rows = [row for row in rows if row["transaction_type"] == "purchase_cancellation"]
        sales_rows = [row for row in rows if row["transaction_type"] == "sales_cancellation"]

        self.assertEqual([(row["source_purchase_date"], row["fifo_cost_amount"]) for row in purchase_rows], [("2026-05-11", -180)])
        self.assertEqual([(row["source_purchase_date"], row["fifo_cost_amount"]) for row in sales_rows], [("2026-05-10", -240), ("2026-05-10", -160)])

    def test_fifo_unallocated_sale_and_cancellation_rows_preserve_quantity_changes(self):
        partial_sales = [
            self._sale_record("2026-01-01", 10, 1, 2),
            self._sale_record("2026-01-02", -4, 2, 3),
        ]
        partial_rows = build_fifo_calculation_rows(self._fifo_report_data([], partial_sales, -6))
        self.assertEqual(
            [
                (row["transaction_type"], row["transaction_quantity"], row["allocation_type"], row["unallocated_quantity"])
                for row in partial_rows
            ],
            [("sales", 10, "unallocated", 10), ("sales_cancellation", -4, "unallocated_cancellation", -4)],
        )

        fully_cancelled_sales = [
            self._sale_record("2026-01-01", 10, 1, 2),
            self._sale_record("2026-01-02", -10, 2, 3),
        ]
        fully_cancelled_rows = build_fifo_calculation_rows(self._fifo_report_data([], fully_cancelled_sales, 0))
        self.assertEqual(
            [
                (row["transaction_type"], row["transaction_quantity"], row["allocation_type"], row["unallocated_quantity"])
                for row in fully_cancelled_rows
            ],
            [("sales", 10, "unallocated", 10), ("sales_cancellation", -10, "unallocated_cancellation", -10)],
        )

    def test_writes_no_formulas_or_external_links_and_uses_native_internal_charts(self):
        sources = copy.deepcopy(load_sources(EXAMPLE_DIR))
        sources["sales"]["records"][0]["item_name"] = "=not_a_formula"
        output, workbook = self._create_workbook(sources)

        self.assertFalse(
            any(
                cell.data_type == "f"
                for worksheet in workbook.worksheets
                for row in worksheet.iter_rows()
                for cell in row
            )
        )
        injected = workbook["매출"]["E2"]
        self.assertEqual(injected.value, "=not_a_formula")
        self.assertEqual(injected.data_type, "s")
        self.assertGreaterEqual(sum(len(worksheet._charts) for worksheet in workbook.worksheets), 1)
        self.assertEqual(workbook._external_links, [])

        with zipfile.ZipFile(output) as archive:
            names = archive.namelist()
            chart_files = [name for name in names if name.startswith("xl/charts/")]
            self.assertTrue(chart_files)
            self.assertFalse(any("externalLinks" in name for name in names))
            self.assertFalse(any("connections" in name for name in names))
            for chart_file in chart_files:
                chart_xml = archive.read(chart_file).decode("utf-8")
                self.assertIn("!$", chart_xml)
                self.assertNotIn("[", chart_xml)

    def test_rate_cells_keep_their_distinct_ratio_and_percentage_point_contracts(self):
        sources = copy.deepcopy(load_sources(EXAMPLE_DIR))
        sources["reconciliation"]["summary"]["gross_profit_status"] = "confirmed"
        output, workbook = self._create_workbook(sources)

        summary_rate = self._labeled_row(workbook, "요약", "매출총이익률")[1]
        profit_rate = self._labeled_row(workbook, "손익분석", "매출총이익률")[1]
        self.assertAlmostEqual(summary_rate.value, 4500 / 9460)
        self.assertEqual(summary_rate.number_format, "0.0%")
        self.assertEqual(profit_rate.number_format, "0.0%")

        weekly_header_row, weekly_rate_column = self._header_column(workbook, "손익분석", "margin_rate")
        weekly_rate = workbook["손익분석"].cell(weekly_header_row + 1, weekly_rate_column)
        self.assertAlmostEqual(weekly_rate.value, 50.285714285714285)
        self.assertEqual(weekly_rate.number_format, "0.0\\%")

        principal_header_row, principal_rate_column = self._header_column(workbook, "매출분석", "margin_rate")
        principal_rate = workbook["매출분석"].cell(principal_header_row + 1, principal_rate_column)
        self.assertAlmostEqual(principal_rate.value, 33.77777777777778)
        self.assertEqual(principal_rate.number_format, "0.0\\%")

        inventory_header_row, inventory_rate_column = self._header_column(workbook, "재고분석", "gross_profit_rate")
        inventory_rate = workbook["재고분석"].cell(inventory_header_row + 1, inventory_rate_column)
        self.assertAlmostEqual(inventory_rate.value, 53.714285714285715)
        self.assertEqual(inventory_rate.number_format, "0.0\\%")
        self.assertTrue(output.exists())

    def test_charts_include_only_the_requested_non_contiguous_series(self):
        output, workbook = self._create_workbook()
        first_chart = workbook["요약"]._charts[0]
        flow_header_row, purchase_column = self._header_column(workbook, "재고분석", "purchase_increase")
        sales_column = next(
            cell.column
            for cell in workbook["재고분석"][flow_header_row]
            if cell.value == "sales_amount"
        )
        purchase_coordinate = workbook["재고분석"].cell(flow_header_row, purchase_column).coordinate
        sales_coordinate = workbook["재고분석"].cell(flow_header_row, sales_column).coordinate

        self.assertEqual(len(first_chart.series), 2)
        title_formulas = [series.tx.strRef.f for series in first_chart.series]
        value_formulas = [series.val.numRef.f for series in first_chart.series]
        self.assertEqual([formula.rsplit("!", 1)[1] for formula in title_formulas], [purchase_coordinate, sales_coordinate])
        self.assertEqual(
            [re.search(r"!\$([A-Z]+)\$", formula).group(1) for formula in value_formulas],
            [getattr(workbook["재고분석"].cell(flow_header_row, column), "column_letter") for column in (purchase_column, sales_column)],
        )
        self.assertTrue(output.exists())

    def test_profit_sheet_matches_html_vat_settlement_values_and_statuses(self):
        output, workbook = self._create_workbook()
        report_summary = build_report_data(load_sources(EXAMPLE_DIR))["summary"]
        vat_row = self._labeled_row(workbook, "손익분석", "잠정 부가세 납부 예상액")
        post_vat_row = self._labeled_row(workbook, "손익분석", "검증 필요")

        self.assertEqual(vat_row[1].value, abs(report_summary["vat_settlement_amount"]))
        self.assertEqual(vat_row[2].value, report_summary["vat_settlement_status"])
        self.assertEqual(post_vat_row[1].value, report_summary["post_vat_reference_amount"])
        self.assertEqual(post_vat_row[2].value, report_summary["post_vat_reference_status"])
        self.assertEqual(self._labeled_row(workbook, "손익분석", "매출 합계금액")[1].value, report_summary["period_sales_total_amount"])
        self.assertEqual(self._labeled_row(workbook, "손익분석", "매입 합계금액")[1].value, report_summary["period_purchase_total_amount"])
        self.assertTrue(output.exists())

    def test_chart_data_shortage_is_a_warning_and_skips_only_unsupported_charts(self):
        sources = copy.deepcopy(load_sources(EXAMPLE_DIR))
        sources["purchase"]["records"] = []
        sources["sales"]["records"] = []
        sources["inventory"]["records"] = []
        output, workbook = self._create_workbook(sources)
        validation_codes = {
            row[1].value
            for row in workbook["검증결과"].iter_rows(min_row=2, max_col=2)
        }
        execution_values = {
            row[0].value: row[1].value
            for row in workbook["실행정보"].iter_rows(min_row=2, max_col=2)
        }

        self.assertEqual(workbook["요약"]["B2"].value, "completed_with_warnings")
        self.assertIn("chart_data_unavailable", validation_codes)
        self.assertIn("chart_data_unavailable", execution_values["적용된 경고"])
        self.assertEqual(len(workbook["요약"]._charts), 1)
        self.assertIn("건수", str(workbook["요약"]._charts[0].x_axis.title))
        self.assertTrue(output.exists())

    def test_missing_optional_metadata_generates_a_warning_without_inventing_records(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_dir = Path(tmpdir) / "input"
            shutil.copytree(EXAMPLE_DIR, input_dir)
            (input_dir / "sales_voucher_metadata.json").unlink()
            output = Path(tmpdir) / "analysis_result.xlsx"

            generate_report(input_dir, output)
            workbook = load_workbook(output, data_only=False)
            execution_values = {
                row[0].value: row[1].value
                for row in workbook["실행정보"].iter_rows(min_row=2, max_col=2)
            }

            self.assertEqual(workbook["요약"]["B2"].value, "completed_with_warnings")
            self.assertEqual(workbook["전표메타"].max_row, 1)
            self.assertEqual(execution_values["제공되지 않은 선택 입력"], "sales_voucher_metadata.json")
            self.assertIn("sales_voucher_metadata_missing", execution_values["적용된 경고"])

    def test_missing_required_json_fails_clearly(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_dir = Path(tmpdir) / "input"
            shutil.copytree(EXAMPLE_DIR, input_dir)
            (input_dir / "sales.json").unlink()

            with self.assertRaisesRegex(FileNotFoundError, "required sales JSON is missing"):
                generate_report(input_dir, Path(tmpdir) / "analysis_result.xlsx")

    def test_missing_core_reconciliation_result_fails_clearly(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_dir = Path(tmpdir) / "input"
            shutil.copytree(EXAMPLE_DIR, input_dir)
            reconciliation_path = input_dir / "inventory_reconciliation.json"
            reconciliation = json.loads(reconciliation_path.read_text(encoding="utf-8"))
            del reconciliation["summary"]["gross_profit"]
            reconciliation_path.write_text(json.dumps(reconciliation), encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "required reconciliation results are missing.*gross_profit"):
                load_analysis_sources(input_dir)

    def test_invalid_core_reconciliation_values_fail_schema_validation(self):
        invalid_values = (
            ("summary", "gross_profit", {}),
            ("summary", "ending_signed_stock_quantity", []),
            ("summary", "fifo_sales_cost_amount", float("nan")),
            ("summary", "gross_profit_status", "not-a-status"),
            ("metadata", "period_start", "2026/05/10"),
        )
        for section, field, invalid_value in invalid_values:
            with self.subTest(field=field), tempfile.TemporaryDirectory() as tmpdir:
                input_dir = Path(tmpdir) / "input"
                shutil.copytree(EXAMPLE_DIR, input_dir)
                reconciliation_path = input_dir / "inventory_reconciliation.json"
                reconciliation = json.loads(reconciliation_path.read_text(encoding="utf-8"))
                reconciliation[section][field] = invalid_value
                reconciliation_path.write_text(json.dumps(reconciliation), encoding="utf-8")

                with self.assertRaisesRegex(ValueError, f"missing or invalid.*{field}"):
                    load_analysis_sources(input_dir)

    def test_clean_validation_result_stays_completed_without_a_validation_chart_warning(self):
        sources = copy.deepcopy(load_sources(EXAMPLE_DIR))
        reconciliation = sources["reconciliation"]
        reconciliation["errors"] = []
        reconciliation["warnings"] = []
        reconciliation["amount_validation_errors"] = []
        reconciliation["summary"].update(
            amount_validation_errors=[],
            amount_validation_error_count=0,
            amount_validation_status="valid",
            gross_profit_status="confirmed",
            profit_status="confirmed",
        )

        output, workbook = self._create_workbook(sources)
        execution_values = {
            row[0].value: row[1].value
            for row in workbook["실행정보"].iter_rows(min_row=2, max_col=2)
        }

        self.assertEqual(workbook["요약"]["B2"].value, "completed")
        self.assertEqual(execution_values["결과 상태"], "completed")
        self.assertEqual(execution_values["적용된 경고"], "없음")
        self.assertEqual(len(workbook["요약"]._charts), 5)
        self.assertTrue(output.exists())

    def test_reopen_failure_does_not_leave_a_partial_final_or_temporary_workbook(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "analysis_result.xlsx"

            with patch("generate_analysis_result_excel.load_workbook", side_effect=OSError("simulated reopen failure")):
                with self.assertRaisesRegex(OSError, "simulated reopen failure"):
                    export_workbook(load_sources(EXAMPLE_DIR), output)

            self.assertFalse(output.exists())
            self.assertEqual(list(Path(tmpdir).glob(".analysis_result.*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
