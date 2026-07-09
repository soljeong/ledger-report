#!/usr/bin/env python3
from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from export_sales_voucher_excel import export_workbook, load_sales


BASE_DIR = Path(__file__).resolve().parent
EXAMPLE_DIR = BASE_DIR / "examples" / "dummy_json"


class ExportSalesVoucherExcelTests(unittest.TestCase):
    def test_exports_sales_grouped_by_date_and_voucher_with_collapsed_detail_rows(self) -> None:
        sales = load_sales(EXAMPLE_DIR / "sales.json")

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "sales_vouchers.xlsx"
            export_workbook(sales, output_path)

            workbook = load_workbook(output_path, data_only=True)
            self.assertEqual(workbook.sheetnames, ["매출_전표별", "전표_메타입력"])

            grouped = workbook["매출_전표별"]
            meta = workbook["전표_메타입력"]

            self.assertEqual(grouped.max_row, 1 + 3 + 4)
            self.assertEqual(meta.max_row, 1 + 3)

            self.assertEqual(grouped["A1"].value, "행유형")
            self.assertEqual(grouped["B1"].value, "전표키")
            self.assertEqual(grouped["P1"].value, "담당자")
            self.assertEqual(grouped["Q1"].value, "확인상태")
            self.assertEqual(grouped["R1"].value, "메모")

            self.assertEqual(grouped["A2"].value, "전표")
            self.assertEqual(grouped["B2"].value, "2026-05-08-1")
            self.assertEqual(grouped["C2"].value, "2026-05-08")
            self.assertEqual(grouped["D2"].value, 1)
            self.assertEqual(grouped["E2"].value, "Sample Robotics")
            self.assertEqual(grouped["F2"].value, 2)
            self.assertEqual(grouped["G2"].value, 60)
            self.assertEqual(grouped["K2"].value, 462000)

            self.assertEqual(grouped["A3"].value, "상세")
            self.assertEqual(grouped["B3"].value, "2026-05-08-1")
            self.assertEqual(grouped["H3"].value, 8000)
            self.assertEqual(grouped["L3"].value, 1001)
            self.assertEqual(grouped.row_dimensions[3].outlineLevel, 1)
            self.assertTrue(grouped.row_dimensions[3].hidden)

            self.assertEqual(meta["A1"].value, "전표키")
            self.assertEqual(meta["A2"].value, "2026-05-08-1")
            self.assertEqual(meta["B2"].value, "2026-05-08")
            self.assertEqual(meta["C2"].value, 1)
            self.assertEqual(meta["D2"].value, "Sample Robotics")
            self.assertEqual(meta["E2"].value, 2)
            self.assertEqual(meta["F2"].value, 60)
            self.assertEqual(meta["I2"].value, 462000)

    def test_rejects_non_sales_payload(self) -> None:
        payload = {"metadata": {"transaction_type": "purchase"}, "records": []}

        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = Path(tmpdir) / "not_sales.json"
            input_path.write_text(json.dumps(payload), encoding="utf-8")

            with self.assertRaisesRegex(ValueError, "sales JSON"):
                load_sales(input_path)


if __name__ == "__main__":
    unittest.main()
