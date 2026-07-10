#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
from collections import OrderedDict
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_DATA_DIR = BASE_DIR / "private_intermediate"
DEFAULT_INPUT = DEFAULT_DATA_DIR / "purchase.json"
DEFAULT_OUTPUT = DEFAULT_DATA_DIR / "purchase_voucher_groups.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
VOUCHER_FILL = PatternFill("solid", fgColor="EAF4E1")
DETAIL_FILL = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT = Font(bold=True)
VOUCHER_FONT = Font(bold=True)

GROUPED_HEADERS = [
    "행유형",
    "전표키",
    "일자",
    "전표",
    "거래처",
    "행수",
    "수량",
    "단가",
    "공급금액",
    "부가세",
    "합계금액",
    "등록번호",
    "품명",
    "규격",
    "원본행",
    "담당자",
    "확인상태",
    "메모",
]

META_HEADERS = [
    "전표키",
    "일자",
    "전표",
    "거래처",
    "행수",
    "수량",
    "공급금액",
    "부가세",
    "합계금액",
    "담당자",
    "확인상태",
    "메모",
    "보충자료",
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export purchase rows grouped by date and voucher.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def load_purchase(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if payload.get("metadata", {}).get("transaction_type") != "purchase":
        raise ValueError(f"expected purchase JSON: {path}")
    return payload


def unique_join(values: list[Any]) -> str:
    seen: list[str] = []
    for value in values:
        if value in (None, ""):
            continue
        text = str(value)
        if text not in seen:
            seen.append(text)
    return "\n".join(seen)


def group_purchase_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: OrderedDict[tuple[str, Any], list[dict[str, Any]]] = OrderedDict()
    for row in sorted(records, key=lambda item: (item["date"], item["voucher"], item["excel_row"])):
        groups.setdefault((row["date"], row["voucher"]), []).append(row)

    result: list[dict[str, Any]] = []
    for (date, voucher), rows in groups.items():
        result.append(
            {
                "key": f"{date}-{voucher}",
                "date": date,
                "voucher": voucher,
                "company": unique_join([row.get("company") for row in rows]),
                "row_count": len(rows),
                "quantity": sum(row.get("quantity") or 0 for row in rows),
                "supply_amount": sum(row.get("supply_amount") or 0 for row in rows),
                "vat": sum(row.get("vat") or 0 for row in rows),
                "total_amount": sum(row.get("total_amount") or 0 for row in rows),
                "rows": rows,
            }
        )
    return result


def setup_sheet(ws, headers: list[str]) -> None:
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def apply_number_format(ws, columns: list[str]) -> None:
    for column in columns:
        for cell in ws[column][1:]:
            cell.number_format = "#,##0"


def auto_size(ws) -> None:
    for column_cells in ws.columns:
        width = 8
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            longest_line = max(len(line) for line in str(value).splitlines())
            width = max(width, min(longest_line + 2, 45))
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def add_status_validation(ws, cell_range: str) -> None:
    validation = DataValidation(type="list", formula1='"미확인,확인,보류,보충필요"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(cell_range)


def export_grouped_sheet(wb: Workbook, groups: list[dict[str, Any]]) -> None:
    ws = wb.active
    ws.title = "매입_전표별"
    setup_sheet(ws, GROUPED_HEADERS)
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.showOutlineSymbols = True

    for group in groups:
        voucher_row = ws.max_row + 1
        ws.append(
            [
                "전표",
                group["key"],
                group["date"],
                group["voucher"],
                group["company"],
                group["row_count"],
                group["quantity"],
                None,
                group["supply_amount"],
                group["vat"],
                group["total_amount"],
                None,
                None,
                None,
                None,
                None,
                None,
                None,
            ]
        )
        for cell in ws[voucher_row]:
            cell.fill = VOUCHER_FILL
            cell.font = VOUCHER_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[voucher_row].collapsed = True

        for detail in group["rows"]:
            detail_row = ws.max_row + 1
            ws.append(
                [
                    "상세",
                    group["key"],
                    detail["date"],
                    detail["voucher"],
                    detail.get("company"),
                    None,
                    detail.get("quantity"),
                    detail.get("unit_price"),
                    detail.get("supply_amount"),
                    detail.get("vat"),
                    detail.get("total_amount"),
                    detail.get("product_id"),
                    detail.get("item_name"),
                    detail.get("specification"),
                    detail.get("excel_row"),
                    None,
                    None,
                    None,
                ]
            )
            ws.row_dimensions[detail_row].outlineLevel = 1
            ws.row_dimensions[detail_row].hidden = True
            for cell in ws[detail_row]:
                cell.fill = DETAIL_FILL
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    add_status_validation(ws, f"Q2:Q{max(ws.max_row, 2)}")
    apply_number_format(ws, ["F", "G", "H", "I", "J", "K", "L", "O"])
    auto_size(ws)


def export_meta_sheet(wb: Workbook, groups: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("전표_메타입력")
    setup_sheet(ws, META_HEADERS)

    for group in groups:
        ws.append(
            [
                group["key"],
                group["date"],
                group["voucher"],
                group["company"],
                group["row_count"],
                group["quantity"],
                group["supply_amount"],
                group["vat"],
                group["total_amount"],
                None,
                None,
                None,
                None,
            ]
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    add_status_validation(ws, f"K2:K{max(ws.max_row, 2)}")
    apply_number_format(ws, ["E", "F", "G", "H", "I"])
    auto_size(ws)


def export_workbook(purchase: dict[str, Any], output_path: Path) -> None:
    groups = group_purchase_rows(purchase["records"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    export_grouped_sheet(wb, groups)
    export_meta_sheet(wb, groups)
    wb.save(output_path)


def main() -> int:
    args = parse_args()
    purchase = load_purchase(args.input)
    export_workbook(purchase, args.output)
    print(f"wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
