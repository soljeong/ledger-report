#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT_DIR = BASE_DIR / "private_sources"
DEFAULT_OUTPUT_DIR = BASE_DIR / "private_intermediate"

FIELD_LABELS = {
    "company": ["업체명", "거래처명"],
    "date": ["일자", "일 자", "매출일자"],
    "voucher": ["전표"],
    "item_name": ["품명", "품        명", "품      명"],
    "specification": ["규격", "규    격", "규   격"],
    "product_id": ["등록번호"],
    "item_category": ["상품구분"],
    "quantity": ["수량", "수 량"],
    "unit_price": ["단가", "단   가"],
    "foreign_unit_price": ["외화단가"],
    "exchange_rate": ["환율"],
    "currency": ["화폐"],
    "supply_amount": ["공급금액"],
    "vat": ["부가세"],
    "total_amount": ["합계금액"],
    "input_date": ["입력일자"],
    "kind": ["구분"],
    "memo": ["비고"],
    "brand": ["브랜드"],
    "unit": ["단위"],
    "stock_quantity": ["재고합계"],
    "box_or_roll_stock": ["박스/롤재고"],
    "store_stock": ["매장재고"],
    "warehouse_stock": ["창고재고"],
    "average_cost": ["평균원가"],
    "latest_purchase_price": ["최종매입가"],
    "sales_price": ["판매가"],
    "storage_location": ["보관장소"],
    "notes": ["기타"],
}

LEDGER_REQUIRED_FIELDS = [
    "date",
    "item_name",
    "specification",
    "quantity",
    "unit_price",
    "supply_amount",
    "vat",
    "total_amount",
]
INVENTORY_REQUIRED_FIELDS = ["item_name", "specification", "product_id", "stock_quantity"]

TITLE_RANGE_RE = re.compile(
    r"(?P<start>\d{4}\.\d{2}\.\d{2})\s*~\s*(?P<end>\d{4}\.\d{2}\.\d{2})\s*(?P<label>.+)"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Parse purchase, sales, and inventory Excel workbooks to JSON.")
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=DEFAULT_INPUT_DIR,
        help="Directory containing copied .xlsx files.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help="Directory where JSON output files are written.",
    )
    return parser.parse_args()


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip())


LABEL_TO_FIELD = {
    normalize_header(label): field
    for field, labels in FIELD_LABELS.items()
    for label in labels
}


def parse_title(title: Any) -> dict[str, Any]:
    if not isinstance(title, str):
        return {"raw": title}
    match = TITLE_RANGE_RE.fullmatch(title.strip())
    if not match:
        return {"raw": title}
    return {
        "raw": title,
        "start_date": parse_dot_date(match.group("start")).isoformat(),
        "end_date": parse_dot_date(match.group("end")).isoformat(),
        "label": match.group("label").strip(),
    }


def parse_dot_date(value: str) -> date:
    return datetime.strptime(value, "%Y.%m.%d").date()


def parse_workbook_date(value: Any) -> tuple[str, Any]:
    if isinstance(value, datetime):
        return value.date().isoformat(), value.isoformat()
    if isinstance(value, date):
        return value.isoformat(), value.isoformat()
    if isinstance(value, int):
        raw = value
        text = f"{value:08d}"
    elif isinstance(value, float) and value.is_integer():
        raw = int(value)
        text = f"{int(value):08d}"
    elif isinstance(value, str):
        raw = value
        text = re.sub(r"\D", "", value.strip())
    else:
        raise ValueError(f"unsupported date value: {value!r}")

    if not re.fullmatch(r"\d{8}", text):
        raise ValueError(f"date must be YYYYMMDD: {value!r}")
    parsed = datetime.strptime(text, "%Y%m%d").date()
    return parsed.isoformat(), raw


def number_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if value.is_integer():
            return int(value)
        return float(Decimal(str(value)))
    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return int(value)
        return float(value)
    if isinstance(value, str):
        stripped = value.strip().replace(",", "")
        if stripped == "":
            return None
        try:
            parsed = Decimal(stripped)
        except InvalidOperation:
            return value.strip()
        if parsed == parsed.to_integral_value():
            return int(parsed)
        return float(parsed)
    return value


def decimal_number(value: Any) -> Decimal:
    parsed = number_value(value)
    if parsed is None:
        return Decimal("0")
    if isinstance(parsed, bool):
        return Decimal(int(parsed))
    if isinstance(parsed, int):
        return Decimal(parsed)
    if isinstance(parsed, float):
        return Decimal(str(parsed))
    return Decimal(str(parsed))


def text_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def row_to_dict(row_values: tuple[Any, ...], header_map: dict[str, int]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field, index in header_map.items():
        result[field] = row_values[index] if index < len(row_values) else None
    return result


def find_header_row(sheet: Any, required_fields: list[str], max_rows: int = 10) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        header_map: dict[str, int] = {}
        for index, value in enumerate(row):
            field = LABEL_TO_FIELD.get(normalize_header(value))
            if field and field not in header_map:
                header_map[field] = index
        if all(field in header_map for field in required_fields):
            return row_idx, header_map
    raise ValueError(f"could not find header row with fields: {required_fields}")


def classify_workbook(path: Path, sheet_name: str, header_map: dict[str, int]) -> tuple[str, str]:
    text = f"{path.name} {sheet_name}"
    if "stock_quantity" in header_map or "재고" in text:
        return "inventory", "inventory.json"
    if "매입" in text:
        return "purchase", "purchase.json"
    if "매출" in text:
        return "sales", "sales.json"
    raise ValueError(f"could not classify workbook: {path}")


def is_summary_row(values: dict[str, Any]) -> bool:
    raw_date = values.get("date")
    company = values.get("company")
    item_name = values.get("item_name")
    if raw_date == "소계":
        return True
    if company == "합계" and raw_date is None and item_name is None:
        return True
    if raw_date == "합계" and item_name is None:
        return True
    return False


def parse_ledger_workbook(
    path: Path,
    sheet: Any,
    header_row: int,
    header_map: dict[str, int],
    transaction_type: str,
    output_name: str,
) -> tuple[str, dict[str, Any]]:
    records: list[dict[str, Any]] = []
    skipped_summary_rows: list[dict[str, Any]] = []
    skipped_blank_rows = 0
    zero_filled_numeric_rows: list[dict[str, Any]] = []

    for row_idx, row_values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not any(value is not None for value in row_values):
            skipped_blank_rows += 1
            continue
        values = row_to_dict(row_values, header_map)
        if is_summary_row(values):
            skipped_summary_rows.append(
                {
                    "excel_row": row_idx,
                    "company": text_value(values.get("company")),
                    "marker": text_value(values.get("date")),
                    "item_name": text_value(values.get("item_name")),
                    "supply_amount": number_value(values.get("supply_amount")),
                    "vat": number_value(values.get("vat")),
                    "total_amount": number_value(values.get("total_amount")),
                }
            )
            continue

        parsed_date, raw_date = parse_workbook_date(values.get("date"))
        record = {
            "excel_row": row_idx,
            "transaction_type": transaction_type,
            "company": text_value(values.get("company")),
            "date": parsed_date,
            "raw_date": raw_date,
            "voucher": number_value(values.get("voucher")),
            "item_name": text_value(values.get("item_name")),
            "specification": text_value(values.get("specification")),
            "product_id": number_value(values.get("product_id")),
            "item_category": text_value(values.get("item_category")),
            "quantity": number_value(values.get("quantity")),
            "unit_price": number_value(values.get("unit_price")),
            "foreign_unit_price": number_value(values.get("foreign_unit_price")),
            "exchange_rate": number_value(values.get("exchange_rate")),
            "currency": text_value(values.get("currency")),
            "supply_amount": number_value(values.get("supply_amount")),
            "vat": number_value(values.get("vat")),
            "total_amount": number_value(values.get("total_amount")),
            "input_date": parse_workbook_date(values["input_date"])[0] if values.get("input_date") else None,
            "kind": text_value(values.get("kind")),
            "memo": text_value(values.get("memo")),
            "brand": text_value(values.get("brand")),
            "unit": text_value(values.get("unit")),
        }
        missing_fields = [
            field
            for field in ["unit_price", "supply_amount", "vat", "total_amount"]
            if record[field] is None
        ]
        if missing_fields:
            zero_filled_numeric_rows.append({"excel_row": row_idx, "missing_fields": missing_fields})
            for field in missing_fields:
                record[field] = 0
        records.append(record)

    def sum_numeric(field: str) -> int | float:
        return sum(row[field] for row in records if row[field] is not None)

    detail_totals = {
        "quantity": sum_numeric("quantity"),
        "supply_amount": sum_numeric("supply_amount"),
        "vat": sum_numeric("vat"),
        "total_amount": sum_numeric("total_amount"),
    }
    workbook_total = next((row for row in skipped_summary_rows if row["company"] == "합계"), None)
    if workbook_total:
        for key in ["supply_amount", "vat", "total_amount"]:
            expected = detail_totals[key]
            actual = workbook_total[key]
            if actual != expected:
                raise ValueError(
                    f"{path.name} {key} total mismatch: detail sum {expected!r}, workbook total {actual!r}"
                )

    parsed_dates = [datetime.strptime(row["date"], "%Y-%m-%d").date() for row in records]
    payload = {
        "metadata": {
            "source_file": path.name,
            "sheet_name": sheet.title,
            "transaction_type": transaction_type,
            "header_row": header_row,
            "headers": {field: sheet.cell(header_row, index + 1).value for field, index in header_map.items()},
            "record_count": len(records),
            "unique_product_ids": len({row["product_id"] for row in records if row["product_id"] is not None}),
            "skipped_summary_row_count": len(skipped_summary_rows),
            "skipped_blank_row_count": skipped_blank_rows,
            "date_min": min(parsed_dates).isoformat() if parsed_dates else None,
            "date_max": max(parsed_dates).isoformat() if parsed_dates else None,
            "detail_totals": detail_totals,
            "workbook_total_row": workbook_total,
            "skipped_summary_rows": skipped_summary_rows,
            "zero_filled_numeric_rows": zero_filled_numeric_rows,
        },
        "records": records,
    }
    return output_name, payload


def parse_inventory_workbook(
    path: Path,
    sheet: Any,
    header_row: int,
    header_map: dict[str, int],
    output_name: str,
) -> tuple[str, dict[str, Any]]:
    records: list[dict[str, Any]] = []
    skipped_blank_rows = 0
    blank_stock_quantity_rows: list[int] = []
    invalid_stock_quantity_rows: list[dict[str, Any]] = []
    negative_stock_rows: list[dict[str, Any]] = []
    stock_quantity = Decimal("0")
    stock_amount_at_average_cost = Decimal("0")
    stock_amount_at_latest_purchase_price = Decimal("0")

    for row_idx, row_values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not any(value is not None for value in row_values):
            skipped_blank_rows += 1
            continue
        values = row_to_dict(row_values, header_map)
        stock_raw = values.get("stock_quantity")
        stock_value = number_value(stock_raw)
        stock_is_numeric = isinstance(stock_value, (int, float, Decimal)) and not isinstance(stock_value, bool)
        if stock_value is None:
            blank_stock_quantity_rows.append(row_idx)
        elif not stock_is_numeric:
            invalid_stock_quantity_rows.append({"excel_row": row_idx, "stock_quantity": stock_value})
        # An invalid snapshot quantity is retained for analysis-time
        # eligibility validation.  It is not silently turned into zero for
        # either the normalized record or the parser's valid-row totals.
        stock = decimal_number(stock_value) if stock_is_numeric else Decimal("0")
        average_cost = decimal_number(values.get("average_cost"))
        latest_purchase_price = decimal_number(values.get("latest_purchase_price"))

        record = {
            "excel_row": row_idx,
            "transaction_type": "inventory",
            "item_name": text_value(values.get("item_name")),
            "specification": text_value(values.get("specification")),
            "product_id": number_value(values.get("product_id")),
            "stock_quantity": stock_value,
            "box_or_roll_stock": number_value(values.get("box_or_roll_stock")),
            "store_stock": number_value(values.get("store_stock")),
            "warehouse_stock": number_value(values.get("warehouse_stock")),
            "average_cost": number_value(values.get("average_cost")),
            "latest_purchase_price": number_value(values.get("latest_purchase_price")),
            "sales_price": number_value(values.get("sales_price")),
            "brand": text_value(values.get("brand")),
            "item_category": text_value(values.get("item_category")),
            "storage_location": text_value(values.get("storage_location")),
            "unit": text_value(values.get("unit")),
            "notes": text_value(values.get("notes")),
        }
        if stock < 0:
            negative_stock_rows.append(
                {
                    "excel_row": row_idx,
                    "item_name": record["item_name"],
                    "product_id": record["product_id"],
                    "stock_quantity": record["stock_quantity"],
                }
            )
        records.append(record)
        stock_quantity += stock
        stock_amount_at_average_cost += stock * average_cost
        stock_amount_at_latest_purchase_price += stock * latest_purchase_price

    def rounded_int(value: Decimal) -> int:
        return int(value.to_integral_value(rounding="ROUND_HALF_UP"))

    payload = {
        "metadata": {
            "source_file": path.name,
            "sheet_name": sheet.title,
            "transaction_type": "inventory",
            "header_row": header_row,
            "headers": {field: sheet.cell(header_row, index + 1).value for field, index in header_map.items()},
            "record_count": len(records),
            "unique_product_ids": len({row["product_id"] for row in records if row["product_id"] is not None}),
            "skipped_blank_row_count": skipped_blank_rows,
            "blank_stock_quantity_row_count": len(blank_stock_quantity_rows),
            "blank_stock_quantity_rows": blank_stock_quantity_rows,
            "invalid_stock_quantity_row_count": len(invalid_stock_quantity_rows),
            "invalid_stock_quantity_rows": invalid_stock_quantity_rows,
            "negative_stock_row_count": len(negative_stock_rows),
            "negative_stock_rows": negative_stock_rows,
            "detail_totals": {
                "stock_quantity": rounded_int(stock_quantity),
                "stock_amount_at_average_cost": rounded_int(stock_amount_at_average_cost),
                "stock_amount_at_latest_purchase_price": rounded_int(stock_amount_at_latest_purchase_price),
            },
        },
        "records": records,
    }
    return output_name, payload


def parse_workbook(path: Path) -> tuple[str, dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"expected exactly one sheet in {path}, got {workbook.sheetnames}")
    sheet = workbook[workbook.sheetnames[0]]

    try:
        header_row, header_map = find_header_row(sheet, INVENTORY_REQUIRED_FIELDS)
    except ValueError:
        header_row, header_map = find_header_row(sheet, LEDGER_REQUIRED_FIELDS)

    transaction_type, output_name = classify_workbook(path, sheet.title, header_map)
    if transaction_type == "inventory":
        return parse_inventory_workbook(path, sheet, header_row, header_map, output_name)
    return parse_ledger_workbook(path, sheet, header_row, header_map, transaction_type, output_name)


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def choose_existing(input_dir: Path, candidates: list[str]) -> Path:
    for name in candidates:
        path = input_dir / name
        if path.exists():
            return path
    raise FileNotFoundError(f"none of these files exist in {input_dir}: {candidates}")


def workbook_candidates(input_dir: Path) -> list[Path]:
    keywords = ("매입", "매출", "재고", "purchase", "sales", "inventory")
    excluded = ("voucher", "전표")
    return [
        path
        for path in sorted(input_dir.glob("*.xlsx"))
        if any(keyword in path.name.lower() for keyword in keywords)
        and not any(keyword in path.name.lower() for keyword in excluded)
    ]


def main() -> int:
    args = parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    input_files = workbook_candidates(args.input_dir)
    if not input_files:
        raise FileNotFoundError(f"no purchase, sales, or inventory .xlsx files found in {args.input_dir}")

    for input_file in input_files:
        output_name, payload = parse_workbook(input_file)
        output_path = args.output_dir / output_name
        write_json(output_path, payload)
        meta = payload["metadata"]
        if meta["transaction_type"] == "inventory":
            print(
                f"{output_path.name}: records={meta['record_count']} "
                f"unique_product_ids={meta['unique_product_ids']} "
                f"stock_qty={meta['detail_totals']['stock_quantity']} "
                f"avg_stock_amount={meta['detail_totals']['stock_amount_at_average_cost']}"
            )
        else:
            print(
                f"{output_path.name}: records={meta['record_count']} "
                f"summary_rows_skipped={meta['skipped_summary_row_count']} "
                f"date_range={meta['date_min']}..{meta['date_max']} "
                f"total={meta['detail_totals']['total_amount']}"
            )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
