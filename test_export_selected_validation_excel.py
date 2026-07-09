#!/usr/bin/env python3
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from export_selected_validation_excel import export_workbook


class ExportSelectedValidationExcelTests(unittest.TestCase):
    def test_exports_only_selected_validation_sections(self) -> None:
        sample_row = {
            "type": "sales",
            "excel_row": 2,
            "date": "2026-05-08",
            "company": "Sample Robotics",
            "quantity": 40,
            "unit_price": 8000,
            "total_amount": 352000,
        }
        validation = {
            "summary": {
                "same_item_different_specs_count": 1,
                "sales_without_purchase_item_count": 1,
                "sales_without_purchase_item_spec_count": 1,
            },
            "same_item_different_specs": [
                {
                    "item_key": "sample part",
                    "item_names": ["Sample Part"],
                    "side_counts": {"purchase": 1, "sales": 1},
                    "specifications": [
                        {"specifications": ["Spec A"], "rows": [sample_row]},
                        {"specifications": ["Spec B"], "rows": [sample_row]},
                    ],
                }
            ],
            "sales_without_purchase_item": [
                {
                    "item_key": "sales only",
                    "item_names": ["Sales Only"],
                    "sales_rows": 1,
                    "sales_total_amount": 352000,
                    "specifications": ["Spec A"],
                    "rows": [sample_row],
                }
            ],
            "sales_without_purchase_item_spec": [
                {
                    "item_key": "sample part",
                    "spec_key": "spec c",
                    "item_names": ["Sample Part"],
                    "specifications": ["Spec C"],
                    "sales_rows": 1,
                    "sales_total_amount": 352000,
                    "rows": [sample_row],
                }
            ],
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            output_path = Path(tmpdir) / "selected.xlsx"
            export_workbook(validation, output_path)

            workbook = load_workbook(output_path, data_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "요약",
                    "같은품명_다른규격",
                    "매출품명_매입없음",
                    "매출품명규격_매입없음",
                ],
            )
            self.assertEqual(workbook["같은품명_다른규격"].max_row - 1, 1)
            self.assertEqual(workbook["매출품명_매입없음"].max_row - 1, 1)
            self.assertEqual(workbook["매출품명규격_매입없음"].max_row - 1, 1)


if __name__ == "__main__":
    unittest.main()
