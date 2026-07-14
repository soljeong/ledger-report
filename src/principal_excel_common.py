from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import generate_analysis_report_html_core as report_core


PRINCIPAL_VOUCHER_SHEET = "원청별 전표손익"
PRINCIPAL_ITEM_SHEET = "원청별 품목수불"
MISSING_PRINCIPAL = "(원청 없음)"

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
PRINCIPAL_FILL = PatternFill("solid", fgColor="BDD7EE")
GROUP_FILL = PatternFill("solid", fgColor="E2F0D9")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
ADJUSTMENT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FONT = Font(bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))

PRINCIPAL_VOUCHER_COLUMNS = [
    "구분",
    "원청",
    "일자",
    "전표",
    "거래처",
    "품목코드",
    "품명",
    "규격",
    "수량",
    "매출단가",
    "매출 공급가액",
    "매출원가",
    "매출총이익",
    "이익률",
    "거래정보",
]

PRINCIPAL_ITEM_COLUMNS = [
    "구분",
    "원청",
    "거래원청",
    "product_id",
    "item_name",
    "specification",
    "기간매입수량",
    "기간매입금액",
    "기간매출수량",
    "기간매출액",
    "매출원가",
    "매출총이익",
    "이익률",
    "기말재고수량",
    "기말재고금액",
    "거래정보",
    "수량",
    "단가",
    "거래 매출원가",
]

# The last display header intentionally duplicates the period-summary "매출원가"
# column.  Keep a distinct internal key so row values cannot overwrite each other.
PRINCIPAL_ITEM_HEADERS = [
    *(PRINCIPAL_ITEM_COLUMNS[:-1]),
    "매출원가",
]


def _decimal(value: Any) -> Decimal | None:
    return report_core.decimal_amount(value)


def _number(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    return int(value) if value == value.to_integral_value() else float(value)


def _transaction_id(row: dict[str, Any]) -> str:
    value = row.get("transaction_id") or row.get("sale_id")
    if value:
        return str(value)
    return f"{row.get('date')}|{row.get('voucher')}|{row.get('excel_row')}"


def _voucher_key(row: dict[str, Any]) -> str:
    return str(row.get("voucher_key") or f"{row.get('date')}-{row.get('voucher')}")


def _principal_map(report_data: dict[str, Any]) -> dict[str, str]:
    return {
        _voucher_key(row): str(row.get("principal") or MISSING_PRINCIPAL)
        for row in report_data["records"].get("sales_voucher_metadata", [])
    }


def _principal_for_sale(row: dict[str, Any], principal_by_voucher: dict[str, str]) -> str:
    return principal_by_voucher.get(_voucher_key(row), MISSING_PRINCIPAL)


def _product_key(row: dict[str, Any]) -> tuple[str, ...]:
    product_id = row.get("product_id")
    if product_id not in (None, ""):
        return ("id", str(product_id))
    return (
        "label",
        str(row.get("item_name") or ""),
        str(row.get("specification") or ""),
    )


def _sort_text(value: Any) -> tuple[int, str]:
    text = str(value or "")
    return (1 if text == MISSING_PRINCIPAL else 0, text)


def _sort_number(value: Any) -> tuple[int, Decimal | str]:
    number = _decimal(value)
    return (0, number) if number is not None else (1, str(value or ""))


def _published_cost(event: dict[str, Any]) -> Decimal | None:
    for name in ("cost_amount_exact", "cost_amount", "fifo_cost_amount"):
        if name not in event or event.get(name) is None:
            continue
        value = _decimal(event.get(name))
        if value is not None:
            return value
    return None


def _sale_cost_map(report_data: dict[str, Any]) -> dict[str, Decimal | None]:
    reconciliation = report_data["reconciliation"]
    costs: dict[str, Decimal | None] = {}
    for event in reconciliation.get("sales_cost_events", []):
        if not isinstance(event, dict) or event.get("event_type") not in {"sale", "sale_cancellation"}:
            continue
        costs[_transaction_id(event)] = _published_cost(event)
    if not costs:
        for allocation in reconciliation.get("sales_allocations", []):
            if not isinstance(allocation, dict) or not allocation.get("in_analysis_period"):
                continue
            costs[_transaction_id(allocation)] = _published_cost(allocation)

    unconfirmed = {
        _transaction_id(allocation)
        for allocation in reconciliation.get("sales_allocations", [])
        if isinstance(allocation, dict)
        and allocation.get("in_analysis_period")
        and (_decimal(allocation.get("unconfirmed_quantity")) or Decimal("0")) != 0
    }
    for transaction_id in unconfirmed:
        costs[transaction_id] = None
    return costs


def _sales_metrics(records: Iterable[dict[str, Any]], cost_by_transaction: dict[str, Decimal | None]) -> dict[str, Any]:
    rows = list(records)
    quantity = sum((_decimal(row.get("quantity")) or Decimal("0") for row in rows), Decimal("0"))
    sales_amount = sum((_decimal(row.get("supply_amount")) or Decimal("0") for row in rows), Decimal("0"))
    costs = [cost_by_transaction.get(_transaction_id(row)) for row in rows]
    cost_complete = all(cost is not None for cost in costs)
    cost_amount = sum((cost for cost in costs if cost is not None), Decimal("0")) if cost_complete else None
    margin_amount = sales_amount - cost_amount if cost_amount is not None else None
    margin_rate = margin_amount / sales_amount if margin_amount is not None and sales_amount != 0 else None
    return {
        "quantity": _number(quantity),
        "sales_amount": _number(sales_amount),
        "cost_amount": _number(cost_amount),
        "margin_amount": _number(margin_amount),
        "margin_rate": float(margin_rate) if margin_rate is not None else None,
    }


def _unique_join(values: Iterable[Any]) -> str:
    seen: list[str] = []
    for value in values:
        if value in (None, ""):
            continue
        text = str(value)
        if text not in seen:
            seen.append(text)
    return ", ".join(seen)


def _excel_value(value: Any, column: str) -> Any:
    if isinstance(value, Decimal):
        return _number(value)
    normalized = column.lower()
    if isinstance(value, str) and ("date" in normalized or column == "일자") and len(value) == 10:
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    return value


def _write_cell(cell, value: Any, column: str) -> None:
    value = _excel_value(value, column)
    cell.value = value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        cell.data_type = "s"
    normalized = column.lower()
    if ("date" in normalized or column == "일자") and value not in (None, ""):
        cell.number_format = "yyyy-mm-dd"
    elif column == "이익률" and isinstance(value, (int, float)):
        cell.number_format = "0.0%"
    elif any(token in normalized for token in ("quantity", "stock", "수량", "excel_row", "원본행")) and isinstance(value, (int, float)):
        cell.number_format = "#,##0.####"
    elif any(token in normalized for token in ("amount", "price", "cost", "금액", "단가", "원가", "공급가", "이익")) and isinstance(value, (int, float)):
        cell.number_format = "#,##0"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _write_header(
    ws: Worksheet,
    row_number: int,
    columns: list[str],
    headers: list[str] | None = None,
) -> None:
    for column_number, (column, header) in enumerate(zip(columns, headers or columns), start=1):
        cell = ws.cell(row_number, column_number)
        _write_cell(cell, header, column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_row(
    ws: Worksheet,
    row_number: int,
    columns: list[str],
    values: dict[str, Any],
    *,
    fill: PatternFill | None = None,
    bold: bool = False,
    outline_level: int = 0,
    hidden: bool = False,
    collapsed: bool = False,
) -> None:
    for column_number, column in enumerate(columns, start=1):
        cell = ws.cell(row_number, column_number)
        _write_cell(cell, values.get(column), column)
        cell.border = THIN_BORDER
        if fill is not None:
            cell.fill = fill
        if bold:
            cell.font = HEADER_FONT
    dimension = ws.row_dimensions[row_number]
    dimension.outlineLevel = outline_level
    dimension.hidden = hidden
    dimension.collapsed = collapsed


def _auto_size(ws: Worksheet, columns: list[str], headers: list[str] | None = None) -> None:
    for column_number, (column, header) in enumerate(zip(columns, headers or columns), start=1):
        width = max(9, len(header) + 2)
        for row_number in range(1, ws.max_row + 1):
            value = ws.cell(row_number, column_number).value
            if value is not None:
                width = max(width, min(max(len(line) for line in str(value).splitlines()) + 2, 34))
        ws.column_dimensions[get_column_letter(column_number)].width = width


def _set_sheet_options(
    ws: Worksheet,
    columns: list[str],
    header_row: int = 3,
    headers: list[str] | None = None,
) -> None:
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{max(ws.max_row, header_row + 1)}"
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.outlinePr.applyStyles = True
    ws.sheet_view.showGridLines = False
    _auto_size(ws, columns, headers)


def _sales_detail_row(row: dict[str, Any], principal: str, cost_by_transaction: dict[str, Decimal | None]) -> dict[str, Any]:
    sales_amount = _decimal(row.get("supply_amount")) or Decimal("0")
    cost_amount = cost_by_transaction.get(_transaction_id(row))
    margin_amount = sales_amount - cost_amount if cost_amount is not None else None
    margin_rate = margin_amount / sales_amount if margin_amount is not None and sales_amount != 0 else None
    quantity = _decimal(row.get("quantity"))
    return {
        "구분": "매출취소" if quantity is not None and quantity < 0 else "매출",
        "원청": principal,
        "일자": row.get("date"),
        "전표": row.get("voucher"),
        "거래처": row.get("company"),
        "품목코드": row.get("product_id"),
        "품명": row.get("item_name"),
        "규격": row.get("specification"),
        "수량": row.get("quantity"),
        "매출단가": row.get("unit_price"),
        "매출 공급가액": _number(sales_amount),
        "매출원가": _number(cost_amount),
        "매출총이익": _number(margin_amount),
        "이익률": float(margin_rate) if margin_rate is not None else None,
        "거래정보": _transaction_id(row),
    }

