from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from src.principal_excel_common import (
    ADJUSTMENT_FILL,
    GROUP_FILL,
    MISSING_PRINCIPAL,
    PRINCIPAL_FILL,
    PRINCIPAL_ITEM_COLUMNS,
    PRINCIPAL_ITEM_SHEET,
    TOTAL_FILL,
    _decimal,
    _number,
    _principal_for_sale,
    _principal_map,
    _product_key,
    _sale_cost_map,
    _sales_metrics,
    _set_sheet_options,
    _sort_text,
    _transaction_id,
    _write_header,
    _write_row,
)

def _inventory_by_product(report_data: dict[str, Any]) -> dict[tuple[str, ...], dict[str, Any]]:
    return {
        _product_key(row): row
        for row in report_data["analysis"].get("inventory_rows", [])
        if isinstance(row, dict)
    }


def _display_product_rows(report_data: dict[str, Any], fifo_rows: list[dict[str, Any]]) -> dict[tuple[str, ...], dict[str, Any]]:
    display: dict[tuple[str, ...], dict[str, Any]] = {}
    sources = [
        report_data["analysis"].get("inventory_rows", []),
        report_data["records"].get("inventory", []),
        report_data["records"].get("period_sales", []),
        report_data["records"].get("period_purchase", []),
        fifo_rows,
    ]
    for records in sources:
        for row in records:
            if not isinstance(row, dict):
                continue
            key = _product_key(row)
            current = display.setdefault(key, {})
            for field in ("product_id", "item_name", "specification"):
                if current.get(field) in (None, "") and row.get(field) not in (None, ""):
                    current[field] = row.get(field)
    return display


def _purchase_metrics(
    product_key: tuple[str, ...],
    inventory_by_product: dict[tuple[str, ...], dict[str, Any]],
    period_purchases: dict[tuple[str, ...], list[dict[str, Any]]],
) -> tuple[Any, Any]:
    inventory_row = inventory_by_product.get(product_key, {})
    if "period_purchase_quantity" in inventory_row or "period_purchase_cost_amount" in inventory_row:
        return inventory_row.get("period_purchase_quantity"), inventory_row.get("period_purchase_cost_amount")
    records = period_purchases.get(product_key, [])
    quantity = sum((_decimal(row.get("quantity")) or Decimal("0") for row in records), Decimal("0"))
    amount = sum((_decimal(row.get("supply_amount")) or Decimal("0") for row in records), Decimal("0"))
    return _number(quantity), _number(amount)


def _fifo_rows_for_principal(
    rows: list[dict[str, Any]],
    principal: str,
    sales_by_transaction: dict[str, dict[str, Any]],
    principal_by_voucher: dict[str, str],
) -> list[tuple[dict[str, Any], str | None]]:
    selected: list[tuple[dict[str, Any], str | None]] = []
    for row in rows:
        transaction_type = str(row.get("transaction_type") or "")
        if transaction_type.startswith("purchase"):
            selected.append((row, None))
            continue
        sale = sales_by_transaction.get(str(row.get("transaction_id") or ""), {})
        transaction_principal = _principal_for_sale(sale, principal_by_voucher) if sale else MISSING_PRINCIPAL
        if transaction_principal == principal:
            selected.append((row, transaction_principal))
    return selected


def write_principal_item_sheet(
    workbook: Workbook,
    report_data: dict[str, Any],
    fifo_rows: list[dict[str, Any]],
) -> Worksheet:
    ws = workbook.create_sheet(PRINCIPAL_ITEM_SHEET)
    ws.cell(1, 1, "원청별 품목 수불")
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.cell(2, 1, "매출은 해당 원청 건만, 매입과 기말재고는 해당 품목 전체를 표시합니다. 공통 품목의 중복 기말재고금액은 하단에서 차감합니다.")
    _write_header(ws, 3, PRINCIPAL_ITEM_COLUMNS)

    principal_by_voucher = _principal_map(report_data)
    cost_by_transaction = _sale_cost_map(report_data)
    inventory_by_product = _inventory_by_product(report_data)
    display_by_product = _display_product_rows(report_data, fifo_rows)

    sales_by_product_principal: dict[tuple[str, ...], dict[str, list[dict[str, Any]]]] = defaultdict(lambda: defaultdict(list))
    sales_by_transaction: dict[str, dict[str, Any]] = {}
    for row in report_data["records"].get("period_sales", []):
        key = _product_key(row)
        principal = _principal_for_sale(row, principal_by_voucher)
        sales_by_product_principal[key][principal].append(row)
        sales_by_transaction[_transaction_id(row)] = row

    period_purchases: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in report_data["records"].get("period_purchase", []):
        period_purchases[_product_key(row)].append(row)

    fifo_by_product: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in fifo_rows:
        fifo_by_product[_product_key(row)].append(row)

    all_product_keys = set(display_by_product) | set(inventory_by_product) | set(sales_by_product_principal) | set(period_purchases) | set(fifo_by_product)
    principals_by_product: dict[tuple[str, ...], set[str]] = {}
    for key in all_product_keys:
        principals = set(sales_by_product_principal.get(key, {}))
        principals_by_product[key] = principals or {MISSING_PRINCIPAL}

    products_by_principal: dict[str, list[tuple[str, ...]]] = defaultdict(list)
    for key, principals in principals_by_product.items():
        for principal in principals:
            products_by_principal[principal].append(key)

    row_number = 4
    displayed_inventory_quantity = Decimal("0")
    displayed_inventory_amount = Decimal("0")
    duplicate_adjustment = Decimal("0")

    for key, principals in principals_by_product.items():
        inventory_row = inventory_by_product.get(key, {})
        ending_quantity = _decimal(inventory_row.get("ending_signed_stock_quantity")) or Decimal("0")
        ending_amount = _decimal(inventory_row.get("ending_fifo_inventory_amount")) or Decimal("0")
        displayed_inventory_quantity += ending_quantity * len(principals)
        displayed_inventory_amount += ending_amount * len(principals)
        duplicate_adjustment += ending_amount * max(len(principals) - 1, 0)

    def product_sort_key(key: tuple[str, ...]) -> tuple[str, str, str]:
        display = display_by_product.get(key, {})
        return (
            str(display.get("product_id") or ""),
            str(display.get("item_name") or ""),
            str(display.get("specification") or ""),
        )

    for principal in sorted(products_by_principal, key=_sort_text):
        product_keys = sorted(products_by_principal[principal], key=product_sort_key)
        principal_sales = [row for key in product_keys for row in sales_by_product_principal.get(key, {}).get(principal, [])]
        principal_metrics = _sales_metrics(principal_sales, cost_by_transaction)
        principal_purchase_quantity = Decimal("0")
        principal_purchase_amount = Decimal("0")
        principal_inventory_quantity = Decimal("0")
        principal_inventory_amount = Decimal("0")
        for key in product_keys:
            purchase_quantity, purchase_amount = _purchase_metrics(key, inventory_by_product, period_purchases)
            principal_purchase_quantity += _decimal(purchase_quantity) or Decimal("0")
            principal_purchase_amount += _decimal(purchase_amount) or Decimal("0")
            inventory_row = inventory_by_product.get(key, {})
            principal_inventory_quantity += _decimal(inventory_row.get("ending_signed_stock_quantity")) or Decimal("0")
            principal_inventory_amount += _decimal(inventory_row.get("ending_fifo_inventory_amount")) or Decimal("0")

        principal_summary_row = row_number
        _write_row(
            ws,
            row_number,
            PRINCIPAL_ITEM_COLUMNS,
            {
                "구분": "원청 합계",
                "원청": principal,
                "기간매입수량": _number(principal_purchase_quantity),
                "기간매입금액": _number(principal_purchase_amount),
                "기간매출수량": principal_metrics["quantity"],
                "기간매출액": principal_metrics["sales_amount"],
                "매출원가": principal_metrics["cost_amount"],
                "매출총이익": principal_metrics["margin_amount"],
                "이익률": principal_metrics["margin_rate"],
                "기말재고수량": _number(principal_inventory_quantity),
                "기말재고금액": _number(principal_inventory_amount),
            },
            fill=PRINCIPAL_FILL,
            bold=True,
            collapsed=bool(product_keys),
        )
        row_number += 1

        for key in product_keys:
            display = display_by_product.get(key, {})
            inventory_row = inventory_by_product.get(key, {})
            sales_records = sales_by_product_principal.get(key, {}).get(principal, [])
            sales_metrics = _sales_metrics(sales_records, cost_by_transaction)
            purchase_quantity, purchase_amount = _purchase_metrics(key, inventory_by_product, period_purchases)
            item_summary_row = row_number
            item_fifo_rows = _fifo_rows_for_principal(
                fifo_by_product.get(key, []),
                principal,
                sales_by_transaction,
                principal_by_voucher,
            )
            _write_row(
                ws,
                row_number,
                PRINCIPAL_ITEM_COLUMNS,
                {
                    "구분": "품목 합계",
                    "원청": principal,
                    "product_id": display.get("product_id"),
                    "item_name": display.get("item_name"),
                    "specification": display.get("specification"),
                    "기간매입수량": purchase_quantity,
                    "기간매입금액": purchase_amount,
                    "기간매출수량": sales_metrics["quantity"],
                    "기간매출액": sales_metrics["sales_amount"],
                    "매출원가": sales_metrics["cost_amount"],
                    "매출총이익": sales_metrics["margin_amount"],
                    "이익률": sales_metrics["margin_rate"],
                    "기말재고수량": inventory_row.get("ending_signed_stock_quantity"),
                    "기말재고금액": inventory_row.get("ending_fifo_inventory_amount"),
                },
                fill=GROUP_FILL,
                bold=True,
                outline_level=1,
                hidden=True,
                collapsed=bool(item_fifo_rows),
            )
            row_number += 1
            for fifo_row, transaction_principal in item_fifo_rows:
                transaction_type = str(fifo_row.get("transaction_type") or "")
                if transaction_type == "purchase_cancellation":
                    detail_type = "매입취소"
                elif transaction_type.startswith("purchase"):
                    detail_type = "매입"
                elif transaction_type == "sales_cancellation":
                    detail_type = "매출취소"
                else:
                    detail_type = "매출"
                detail = {
                    "구분": detail_type,
                    "원청": principal,
                    "거래원청": transaction_principal,
                    **{column: fifo_row.get(column) for column in PRINCIPAL_ITEM_COLUMNS if column in fifo_row},
                }
                _write_row(
                    ws,
                    row_number,
                    PRINCIPAL_ITEM_COLUMNS,
                    detail,
                    outline_level=2,
                    hidden=True,
                )
                row_number += 1
            ws.row_dimensions[item_summary_row].collapsed = bool(item_fifo_rows)
        ws.row_dimensions[principal_summary_row].collapsed = bool(product_keys)

    row_number += 1
    _write_row(
        ws,
        row_number,
        PRINCIPAL_ITEM_COLUMNS,
        {
            "구분": "원청별 표시 합계",
            "기말재고수량": _number(displayed_inventory_quantity),
            "기말재고금액": _number(displayed_inventory_amount),
        },
        fill=TOTAL_FILL,
        bold=True,
    )
    row_number += 1
    _write_row(
        ws,
        row_number,
        PRINCIPAL_ITEM_COLUMNS,
        {
            "구분": "공통 품목 중복 표시 조정",
            "기말재고금액": _number(-duplicate_adjustment),
        },
        fill=ADJUSTMENT_FILL,
        bold=True,
    )
    row_number += 1
    _write_row(
        ws,
        row_number,
        PRINCIPAL_ITEM_COLUMNS,
        {
            "구분": "조정 후 기말 FIFO 재고금액",
            "기말재고금액": _number(displayed_inventory_amount - duplicate_adjustment),
        },
        fill=TOTAL_FILL,
        bold=True,
    )
    row_number += 1
    _write_row(
        ws,
        row_number,
        PRINCIPAL_ITEM_COLUMNS,
        {
            "구분": "수량 안내",
            "원청": "기말재고수량 표시 합계에는 공통 품목이 원청별로 중복 포함됩니다.",
        },
    )

    _set_sheet_options(ws, PRINCIPAL_ITEM_COLUMNS)
    return ws
