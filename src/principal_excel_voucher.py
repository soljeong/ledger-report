from __future__ import annotations

from collections import defaultdict
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from src.principal_excel_common import (
    GROUP_FILL,
    PRINCIPAL_FILL,
    PRINCIPAL_VOUCHER_COLUMNS,
    PRINCIPAL_VOUCHER_SHEET,
    _principal_for_sale,
    _principal_map,
    _sale_cost_map,
    _sales_detail_row,
    _sales_metrics,
    _set_sheet_options,
    _sort_number,
    _sort_text,
    _transaction_id,
    _unique_join,
    _write_header,
    _write_row,
)

def write_principal_voucher_sheet(workbook: Workbook, report_data: dict[str, Any]) -> Worksheet:
    ws = workbook.create_sheet(PRINCIPAL_VOUCHER_SHEET)
    ws.cell(1, 1, "원청별 전표 손익")
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.cell(2, 1, "매출 공급가액 기준이며, 원청 합계 → 전표 합계 → 개별 매출 순서로 펼쳐 볼 수 있습니다.")
    _write_header(ws, 3, PRINCIPAL_VOUCHER_COLUMNS)

    principal_by_voucher = _principal_map(report_data)
    cost_by_transaction = _sale_cost_map(report_data)
    grouped: dict[str, dict[tuple[str, str], list[dict[str, Any]]]] = defaultdict(lambda: defaultdict(list))
    for row in report_data["records"].get("period_sales", []):
        principal = _principal_for_sale(row, principal_by_voucher)
        grouped[principal][(str(row.get("date") or ""), str(row.get("voucher") or ""))].append(row)

    row_number = 4
    for principal in sorted(grouped, key=_sort_text):
        vouchers = grouped[principal]
        principal_records = [row for records in vouchers.values() for row in records]
        principal_metrics = _sales_metrics(principal_records, cost_by_transaction)
        principal_summary_row = row_number
        _write_row(
            ws,
            row_number,
            PRINCIPAL_VOUCHER_COLUMNS,
            {
                "구분": "원청 합계",
                "원청": principal,
                "전표": len(vouchers),
                "거래처": _unique_join(row.get("company") for row in principal_records),
                "수량": principal_metrics["quantity"],
                "매출 공급가액": principal_metrics["sales_amount"],
                "매출원가": principal_metrics["cost_amount"],
                "매출총이익": principal_metrics["margin_amount"],
                "이익률": principal_metrics["margin_rate"],
            },
            fill=PRINCIPAL_FILL,
            bold=True,
            collapsed=bool(vouchers),
        )
        row_number += 1
        for voucher_key in sorted(vouchers, key=lambda value: (value[0], _sort_number(value[1]))):
            records = sorted(vouchers[voucher_key], key=lambda row: (_sort_number(row.get("excel_row")), _transaction_id(row)))
            voucher_metrics = _sales_metrics(records, cost_by_transaction)
            voucher_summary_row = row_number
            _write_row(
                ws,
                row_number,
                PRINCIPAL_VOUCHER_COLUMNS,
                {
                    "구분": "전표 합계",
                    "원청": principal,
                    "일자": records[0].get("date"),
                    "전표": records[0].get("voucher"),
                    "거래처": _unique_join(row.get("company") for row in records),
                    "수량": voucher_metrics["quantity"],
                    "매출 공급가액": voucher_metrics["sales_amount"],
                    "매출원가": voucher_metrics["cost_amount"],
                    "매출총이익": voucher_metrics["margin_amount"],
                    "이익률": voucher_metrics["margin_rate"],
                },
                fill=GROUP_FILL,
                bold=True,
                outline_level=1,
                hidden=True,
                collapsed=True,
            )
            row_number += 1
            for record in records:
                _write_row(
                    ws,
                    row_number,
                    PRINCIPAL_VOUCHER_COLUMNS,
                    _sales_detail_row(record, principal, cost_by_transaction),
                    outline_level=2,
                    hidden=True,
                )
                row_number += 1
            ws.row_dimensions[voucher_summary_row].collapsed = True
        ws.row_dimensions[principal_summary_row].collapsed = True

    if not grouped:
        _write_row(ws, row_number, PRINCIPAL_VOUCHER_COLUMNS, {"구분": "표시할 분석기간 매출이 없습니다."})
    _set_sheet_options(ws, PRINCIPAL_VOUCHER_COLUMNS)
    return ws


