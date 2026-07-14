from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from generate_analysis_report_html_core import load_sources
from generate_analysis_result_excel import SHEET_NAMES, export_workbook
from src.principal_excel_sheets import PRINCIPAL_ITEM_SHEET, PRINCIPAL_VOUCHER_SHEET


BASE_DIR = Path(__file__).resolve().parent
EXAMPLE_DIR = BASE_DIR / "examples" / "dummy_json"


class PrincipalExcelSheetsTests(unittest.TestCase):
    def _create_workbook(self):
        tempdir = tempfile.TemporaryDirectory()
        output = Path(tempdir.name) / "analysis_result.xlsx"
        export_workbook(load_sources(EXAMPLE_DIR), output)
        self.addCleanup(tempdir.cleanup)
        return load_workbook(output, data_only=False)

    @staticmethod
    def _headers(worksheet):
        return {cell.value: cell.column for cell in worksheet[3]}

    @staticmethod
    def _matching_rows(worksheet, headers, **criteria):
        matches = []
        for row_number in range(4, worksheet.max_row + 1):
            if all(worksheet.cell(row_number, headers[column]).value == value for column, value in criteria.items()):
                matches.append(row_number)
        return matches

    def test_principal_voucher_sheet_groups_vouchers_and_keeps_sales_and_cost_on_each_row(self):
        workbook = self._create_workbook()
        self.assertEqual(workbook.sheetnames, SHEET_NAMES)
        self.assertEqual(workbook.sheetnames.index(PRINCIPAL_VOUCHER_SHEET), workbook.sheetnames.index("매출분석") + 1)
        self.assertEqual(workbook.sheetnames.index(PRINCIPAL_ITEM_SHEET), workbook.sheetnames.index(PRINCIPAL_VOUCHER_SHEET) + 1)

        worksheet = workbook[PRINCIPAL_VOUCHER_SHEET]
        headers = self._headers(worksheet)
        self.assertEqual(headers["매출원가"], headers["매출 공급가액"] + 1)
        detail_row = self._matching_rows(worksheet, headers, transaction_id="2026-05-10|2|3")[0]
        self.assertEqual(worksheet.cell(detail_row, headers["구분"]).value, "매출")
        self.assertEqual(worksheet.cell(detail_row, headers["원청"]).value, "Demo Buyer A")
        self.assertEqual(worksheet.cell(detail_row, headers["매출 공급가액"]).value, 2000)
        self.assertEqual(worksheet.cell(detail_row, headers["매출원가"]).value, 840)
        self.assertEqual(worksheet.cell(detail_row, headers["매출총이익"]).value, 1160)
        self.assertAlmostEqual(worksheet.cell(detail_row, headers["이익률"]).value, 0.58)
        self.assertEqual(worksheet.cell(detail_row, headers["이익률"]).number_format, "0.0%")
        self.assertEqual(worksheet.row_dimensions[detail_row].outlineLevel, 2)
        self.assertTrue(worksheet.row_dimensions[detail_row].hidden)

        voucher_row = self._matching_rows(
            worksheet,
            headers,
            구분="전표 합계",
            원청="Demo Buyer A",
            일자=date(2026, 5, 10),
            전표=2,
        )[0]
        self.assertEqual(worksheet.cell(voucher_row, headers["매출 공급가액"]).value, 2000)
        self.assertEqual(worksheet.cell(voucher_row, headers["매출원가"]).value, 840)
        self.assertEqual(worksheet.row_dimensions[voucher_row].outlineLevel, 1)
        self.assertTrue(worksheet.row_dimensions[voucher_row].hidden)
        self.assertTrue(worksheet.row_dimensions[voucher_row].collapsed)

        self.assertFalse(
            any(
                cell.data_type == "f"
                for sheet_name in (PRINCIPAL_VOUCHER_SHEET, PRINCIPAL_ITEM_SHEET)
                for row in workbook[sheet_name].iter_rows()
                for cell in row
            )
        )

    def test_principal_item_sheet_repeats_common_inventory_and_offsets_only_duplicate_amount(self):
        workbook = self._create_workbook()
        worksheet = workbook[PRINCIPAL_ITEM_SHEET]
        headers = self._headers(worksheet)

        item_rows = [
            row_number
            for row_number in self._matching_rows(worksheet, headers, 구분="품목 합계", product_id=1001)
            if worksheet.cell(row_number, headers["원청"]).value in {"Demo Buyer A", "Demo Buyer B"}
        ]
        self.assertEqual(len(item_rows), 2)
        for row_number in item_rows:
            self.assertEqual(worksheet.cell(row_number, headers["기간매입수량"]).value, 11)
            self.assertEqual(worksheet.cell(row_number, headers["기간매입금액"]).value, 1470)
            self.assertEqual(worksheet.cell(row_number, headers["기말재고수량"]).value, 3)
            self.assertEqual(worksheet.cell(row_number, headers["기말재고금액"]).value, 450)
            self.assertEqual(worksheet.row_dimensions[row_number].outlineLevel, 1)
            self.assertTrue(worksheet.row_dimensions[row_number].hidden)

        repeated_purchase_rows = self._matching_rows(
            worksheet,
            headers,
            구분="매입",
            transaction_id="2026-05-01|1|2",
        )
        self.assertEqual(len(repeated_purchase_rows), 2)
        self.assertEqual(
            {worksheet.cell(row_number, headers["원청"]).value for row_number in repeated_purchase_rows},
            {"Demo Buyer A", "Demo Buyer B"},
        )

        buyer_a_sale = self._matching_rows(
            worksheet,
            headers,
            구분="매출",
            원청="Demo Buyer A",
            transaction_id="2026-05-10|2|3",
        )
        buyer_b_sale = self._matching_rows(
            worksheet,
            headers,
            구분="매출",
            원청="Demo Buyer B",
            transaction_id="2026-05-20|3|4",
        )
        self.assertTrue(buyer_a_sale)
        self.assertTrue(buyer_b_sale)
        self.assertFalse(
            self._matching_rows(
                worksheet,
                headers,
                구분="매출",
                원청="Demo Buyer B",
                transaction_id="2026-05-10|2|3",
            )
        )

        adjustment_row = self._matching_rows(worksheet, headers, 구분="공통 품목 중복 표시 조정")[0]
        adjusted_total_row = self._matching_rows(worksheet, headers, 구분="조정 후 기말 FIFO 재고금액")[0]
        self.assertIsNone(worksheet.cell(adjustment_row, headers["기말재고수량"]).value)
        self.assertEqual(worksheet.cell(adjustment_row, headers["기말재고금액"]).value, -450)
        self.assertEqual(worksheet.cell(adjusted_total_row, headers["기말재고금액"]).value, 1510)


if __name__ == "__main__":
    unittest.main()
